import openpyxl,json,datetime
w=openpyxl.load_workbook('/tmp/horsemen-history.xlsx',data_only=True)
f=openpyxl.load_workbook('/tmp/horsemen-history.xlsx',data_only=False)
weeks=[];checks={}
for year in ['2022','2023','2024','2025']:
 s=w[year];cols=[3,5,7,9,11] if year=='2022' else list(range(3,8))
 for r in range(3,24):
  d=s.cell(r,2).value
  if not isinstance(d,datetime.datetime) or not s.cell(r,3).value:continue
  picks=[]
  for c in cols:
   cell=s.cell(r,c);flag=s.cell(r,c+1).value if year=='2022' else w['Backend '+year[2:]].cell(r+1,c).value
   assert flag in [0,1],(year,r,c,flag)
   result='loss' if flag else 'win'
   if '**N/A**'==cell.value:result='absent'
   elif cell.fill.fgColor.rgb=='FFCCCCCC':result='push'
   picks.append({'text':cell.value.strip(),'result':result})
  buy=float(s.cell(r,13 if year=='2022' else 8).value) if year!='2025' else float(w['Backend 25'].cell(r+1,8).value)
  payout=float(s.cell(r,14 if year=='2022' else 8 if year=='2025' else 9).value or 0)
  weeks.append({'date':d.date().isoformat(),'buyIn':buy,'payout':payout,'locked':True,'picks':picks})
 yearly=[x for x in weeks if x['date'].startswith(year)]; totals=[]
 for i in range(5):
  value=0
  for x in yearly:
   losers=sum(p['result']=='loss' for p in x['picks']);people=sum(p['result']!='absent' for p in x['picks'])
   value+=x['payout']-(x['buyIn']*people/losers if losers and x['picks'][i]['result']=='loss' else 0)
  cached=s.cell(23,cols[i]).value
  assert abs(value-cached)<.00001,(year,i,value,cached)
  totals.append(round(value,2))
 checks[year]={'weeks':len(yearly),'balances':totals}
open('lib/history.json','w').write(json.dumps(weeks,indent=2)+'\n')
open('lib/history-audit.json','w').write(json.dumps(checks,indent=2)+'\n')
print(json.dumps(checks,indent=2))
